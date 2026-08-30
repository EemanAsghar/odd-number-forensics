"""Convert writeup.md to writeup.docx with figures embedded. Local only."""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

FIGS = {"figure_violation_rates.png": "Figure 1. DeepSeek R1 instruction-violation rate by condition (n=30 per cell, 95% Wilson CIs). The other four models were 0% in every cell except Llama 3.3 (10% in D2).",
        "figure_dose_response.png": "Figure 2. Authority dose-response for DeepSeek R1: even the hedged 'may reflect' preface produces most of the effect."}

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(1); s.top_margin = s.bottom_margin = Inches(1)

def add_runs(par, text):
    """Handle **bold**, *italic*, `code` inline."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not tok: continue
        if tok.startswith("**"): r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`"): r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif tok.startswith("*"): r = par.add_run(tok[1:-1]); r.italic = True
        else: par.add_run(tok)

def add_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows if not re.match(r"^\s*\|?\s*-", r)]
    t = doc.add_table(rows=len(cells), cols=len(cells[0])); t.style = "Light Grid Accent 1"
    for i, row in enumerate(cells):
        for j, val in enumerate(row):
            if j >= len(t.columns): continue
            p = t.cell(i, j).paragraphs[0]; add_runs(p, val)
            for r in p.runs:
                r.font.size = Pt(9)
                if i == 0: r.bold = True
    doc.add_paragraph()

lines = open("writeup.md").read().split("\n")
i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("```"):
        block = []; i += 1
        while i < len(lines) and not lines[i].startswith("```"): block.append(lines[i]); i += 1
        p = doc.add_paragraph(); r = p.add_run("\n".join(block)); r.font.name = "Consolas"; r.font.size = Pt(9)
        p.paragraph_format.left_indent = Inches(0.3)
    elif ln.startswith("|"):
        block = []
        while i < len(lines) and lines[i].startswith("|"): block.append(lines[i]); i += 1
        add_table(block); continue
    elif ln.startswith("# "):
        h = doc.add_heading(ln[2:], 0)
    elif ln.startswith("## "):
        doc.add_heading(ln[3:], 1)
    elif ln.startswith("### "):
        doc.add_heading(ln[4:], 2)
    elif ln.strip() == "---":
        pass
    elif ln.startswith("> "):
        p = doc.add_paragraph(); add_runs(p, ln[2:]); p.paragraph_format.left_indent = Inches(0.4)
        for r in p.runs: r.italic = True; r.font.color.rgb = RGBColor(0x40, 0x40, 0x50)
    elif re.match(r"^\s*- ", ln):
        p = doc.add_paragraph(style="List Bullet"); add_runs(p, re.sub(r"^\s*- ", "", ln))
    elif re.match(r"^\d+\. ", ln):
        p = doc.add_paragraph(style="List Number"); add_runs(p, re.sub(r"^\d+\. ", "", ln))
    elif "[Insert figure" in ln:
        fname = re.search(r"(figure_\w+\.png)", ln).group(1)
        doc.add_picture(fname, width=Inches(6.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph(); r = cap.add_run(FIGS[fname]); r.italic = True; r.font.size = Pt(9.5)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif ln.strip():
        p = doc.add_paragraph(); add_runs(p, ln)
        if ln.startswith("*") and ln.endswith("*"): p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    i += 1

doc.save("writeup.docx"); print("wrote writeup.docx")
